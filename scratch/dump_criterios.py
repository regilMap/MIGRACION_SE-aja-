import openpyxl

wb = openpyxl.load_workbook('Criterios_Evaluacion_SISMAP.xlsx')
ws = wb.active

print(f"Total rows: {ws.max_row}")
print()

# Print ALL rows grouped by evidencia to understand full structure
prev_ev_id = None
for row in ws.iter_rows(min_row=2, values_only=True):
    ambito, indicador, subindicador, sub_cod, ev_id, ev_cod, ev_nombre, ev_pdf, criterio = row
    if ev_id != prev_ev_id:
        print(f"\n--- EvidenciaID={ev_id} | Cod={ev_cod} | SubInd={sub_cod} | {ev_nombre}")
        prev_ev_id = ev_id
    print(f"  [{criterio}]")
