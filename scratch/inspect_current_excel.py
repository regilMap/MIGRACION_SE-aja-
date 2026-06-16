import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r"c:\Users\regil.batista\Desktop\MIGRACION_SE(aja)\Criterios_Evaluacion_SISMAP.xlsx"
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

print("Unique subindicator codes in current Excel:")
codes = set()
for r in range(2, ws.max_row + 1):
    val = ws.cell(r, 4).value
    if val:
        codes.add(val)

print(sorted(list(codes)))

print("\nDetail of rows in current Excel:")
for r in range(2, ws.max_row + 1):
    sub = ws.cell(r, 4).value
    ev_code = ws.cell(r, 6).value
    crit = ws.cell(r, 9).value
    if sub == "2.03" or "2.03" in str(sub) or "2.03" in str(ev_code):
        print(f"Row {r} | Sub: {sub} | EvCode: {ev_code} | Crit: {crit[:100]}")
