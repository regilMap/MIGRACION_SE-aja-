import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

excel_path = r"c:\Users\regil.batista\Desktop\MIGRACION_SE(aja)\Criterios_Evaluacion_SISMAP.xlsx"

print("--- Verifying Output Excel Sheet ---")
if not os.path.exists(excel_path):
    print("ERROR: Excel file does not exist!")
    sys.exit(1)

wb = openpyxl.load_workbook(excel_path)
ws = wb.active

print(f"Sheet name: {ws.title}")
print(f"Max row count (including header): {ws.max_row}")
print(f"Max col count: {ws.max_column}")

if ws.max_row != 259:
    print(f"ERROR: Expected exactly 259 rows (1 header + 258 data), got {ws.max_row}!")
    sys.exit(1)
else:
    print("SUCCESS: Row count matches expected (259 rows).")

# Verify columns
expected_headers = [
    "Ámbito",
    "Indicador",
    "Subindicador",
    "Subindicador Código",
    "Evidencia ID",
    "Evidencia Código",
    "Nombre de Evidencia (Base de Datos)",
    "Evidencia en PDF",
    "Criterio de Verificación (PDF)"
]

actual_headers = [ws.cell(1, col_idx).value for col_idx in range(1, 10)]
if actual_headers != expected_headers:
    print(f"ERROR: Header row mismatch!\nExpected: {expected_headers}\nGot: {actual_headers}")
    sys.exit(1)
else:
    print("SUCCESS: Headers match perfectly.")

# Verify no missing data in crucial cells
empty_id_count = 0
empty_code_count = 0
empty_crit_count = 0

for row_idx in range(2, 260):
    ev_id = ws.cell(row_idx, 5).value
    ev_code = ws.cell(row_idx, 6).value
    crit = ws.cell(row_idx, 9).value
    
    if not ev_id:
        empty_id_count += 1
    if not ev_code:
        empty_code_count += 1
    if not crit:
        empty_crit_count += 1

print(f"Empty Evidencia IDs: {empty_id_count}")
print(f"Empty Evidencia Codes: {empty_code_count}")
print(f"Empty Criteria: {empty_crit_count}")

if empty_id_count > 0 or empty_code_count > 0 or empty_crit_count > 0:
    print("ERROR: Missing database mappings or empty criteria found!")
    sys.exit(1)
else:
    print("SUCCESS: No missing database codes, IDs, or criteria cells.")
    
# Verify styles
header_cell = ws.cell(1, 1)
print(f"Header fill color: {header_cell.fill.start_color.rgb}")
print(f"Header font: {header_cell.font.name}, Bold={header_cell.font.bold}, Color={header_cell.font.color.rgb}")

data_cell_1 = ws.cell(2, 1)
data_cell_2 = ws.cell(3, 1)
print(f"Row 2 fill color (zebra): {data_cell_1.fill.start_color.rgb}")
print(f"Row 3 fill color (white): {data_cell_2.fill.start_color.rgb}")
print(f"Data font: {data_cell_1.font.name}, Bold={data_cell_1.font.bold}")

print("\nAll verification checks passed successfully!")
